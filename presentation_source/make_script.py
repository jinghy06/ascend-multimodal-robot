from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

build=Path(__file__).resolve().parent
out=build.parent/'开题报告'
c=json.loads((build/'content.json').read_text(encoding='utf-8'))
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.25
normal.paragraph_format.space_after=Pt(6)
doc.styles['Heading 1'].font.size=Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(13)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
title=doc.add_paragraph(style='Title');title.paragraph_format.space_after=Pt(10)
run=title.add_run('开题报告逐页讲稿');run.font.size=Pt(22);run.bold=True
doc.add_paragraph(c['title']).paragraph_format.space_after=Pt(8)
doc.add_paragraph('主控平台：Huawei Atlas 200I DK A2')
doc.add_paragraph('时长安排：正文约 4 分 30 秒，翻页与停顿预留约 10 秒。建议按表中时间排练，控制在 5 分钟内。')
doc.add_paragraph('正文按 PPT 的 10 页顺序编排，可直接朗读。后附答辩备问，不计入汇报时间。')
elapsed=0
md=[f"# 开题报告逐页讲稿\n\n{c['title']}\n\n主控平台：{c['platform']}\n\n正文约 4 分 30 秒，翻页与停顿预留约 10 秒。建议排练控制在 5 分钟内。\n"]
for i,s in enumerate(c['slides'],1):
    if i in [4,7]:doc.add_page_break()
    end=elapsed+s['seconds']
    fmt=lambda x:f'{x//60}:{x%60:02d}'
    h=f"第 {i:02d} 页  {s['title']}"
    doc.add_heading(h,level=1)
    timing=f"建议 {s['seconds']} 秒  /  累计 {fmt(elapsed)}—{fmt(end)}"
    p=doc.add_paragraph(timing);p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=True
    for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('5C7084')
    doc.add_paragraph(s['speech']).paragraph_format.keep_together=True
    md += [f"\n## {h}\n\n{timing}\n\n{s['speech']}\n"]
    elapsed=end

doc.add_page_break()
doc.add_heading('答辩备问',level=1)
doc.add_paragraph('以下内容供问答环节参考，不在 5 分钟汇报正文中朗读。')
md+=['\n# 答辩备问\n\n以下内容不计入汇报时间。\n']
for q,a in c['qa']:
    doc.add_heading(q,level=2);doc.add_paragraph(a)
    md += [f'\n## {q}\n\n{a}\n']
doc.core_properties.title=c['title']+'：开题报告讲稿'
doc.core_properties.subject='逐页讲稿与答辩备问'
doc.core_properties.author='项目组'
footer=sec.footer.paragraphs[0];footer.alignment=2
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for e in list(doc.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
for e in list(doc.styles.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
doc.save(out/'昇腾机器人_开题讲稿.docx')
(out/'昇腾机器人_开题讲稿.md').write_text(''.join(md),encoding='utf-8')
spoken=''.join(s['speech'] for s in c['slides'])
print(json.dumps({'spoken_cjk_chars':len(re.findall('[\u4e00-\u9fff]',spoken)),'total_chars':len(spoken),'planned_seconds':elapsed},ensure_ascii=False))
